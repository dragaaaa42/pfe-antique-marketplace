#!/usr/bin/env python3
"""
Convert the Markdown report to DOCX with premium academic styling.
Navy/Blue/Gold palette matching the PDF design.
"""
import os
import re
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ── Colors ──────────────────────────────────────────────────────
DARK_NAVY = RGBColor(10, 25, 47)
ROYAL_BLUE = RGBColor(30, 58, 138)
SOFT_GOLD = RGBColor(212, 175, 55)
LIGHT_GREY = RGBColor(243, 244, 246)
LIGHT_BLUE = RGBColor(239, 246, 255)
WHITE = RGBColor(255, 255, 255)
BODY_TEXT = RGBColor(31, 41, 55)
MID_GREY = RGBColor(107, 114, 128)
BORDER_CLR = "BFC8E0"

def set_cell_background(cell, fill_hex, color=None, val="clear"):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill_hex)
    shd.set(qn('w:val'), val)
    if color:
        shd.set(qn('w:color'), color)
    tcPr.append(shd)

def set_cell_borders(cell, color="BFC8E0", sz="4"):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right'):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), sz)
        el.set(qn('w:space'), '0')
        el.set(qn('w:color'), color)
        tcBorders.append(el)
    tcPr.append(tcBorders)

def set_cell_padding(cell, top=40, bottom=40, left=60, right=60):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for name, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        el = OxmlElement(f'w:{name}')
        el.set(qn('w:w'), str(val))
        el.set(qn('w:type'), 'dxa')
        tcMar.append(el)
    tcPr.append(tcMar)

def add_page_number(run):
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = "PAGE"
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'separate')
    fldChar3 = OxmlElement('w:fldChar')
    fldChar3.set(qn('w:fldCharType'), 'end')
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)
    run._r.append(fldChar3)

def parse_inline(paragraph, text):
    parts = re.split(r'(\*\*.*?\*\*|\*.*?\*)', text)
    for p in parts:
        if p.startswith('**') and p.endswith('**'):
            run = paragraph.add_run(p[2:-2])
            run.bold = True
        elif p.startswith('*') and p.endswith('*'):
            run = paragraph.add_run(p[1:-1])
            run.italic = True
        else:
            paragraph.add_run(p)

def create_screenshot_placeholder(doc, fig_label, caption):
    """Create a professional framed screenshot placeholder."""
    # Outer table for the frame
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_background(cell, "F8FAFC")
    set_cell_borders(cell, color="94A3B8", sz="6")
    set_cell_padding(cell, top=120, bottom=120, left=100, right=100)

    # Make the row tall enough for a screenshot
    tr = table.rows[0]._tr
    trPr = tr.get_or_add_trPr()
    trHeight = OxmlElement('w:trHeight')
    trHeight.set(qn('w:val'), "3600")  # ~2.5 inches
    trHeight.set(qn('w:hRule'), "atLeast")
    trPr.append(trHeight)

    # Content inside
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.space_before = Pt(20)

    run = p.add_run("\n\nCapture d'ecran a inserer\n")
    run.font.color.rgb = RGBColor(148, 163, 184)
    run.font.size = Pt(13)
    run.font.italic = True

    run2 = p.add_run(f"\n{fig_label}")
    run2.font.color.rgb = MID_GREY
    run2.font.size = Pt(10)

    # Caption below
    caption_p = doc.add_paragraph()
    caption_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_p.space_before = Pt(4)
    caption_p.space_after = Pt(8)
    run = caption_p.add_run(f"{fig_label} : {caption}")
    run.font.color.rgb = ROYAL_BLUE
    run.font.size = Pt(9.5)
    run.font.italic = True


def main():
    doc = Document()

    # ── Configure styles ────────────────────────────────────────
    styles = doc.styles

    h1 = styles['Heading 1']
    h1.font.name = 'Arial'
    h1.font.size = Pt(20)
    h1.font.color.rgb = ROYAL_BLUE
    h1.font.bold = True
    h1_pf = h1.paragraph_format
    h1_pf.space_before = Pt(18)
    h1_pf.space_after = Pt(10)

    h2 = styles['Heading 2']
    h2.font.name = 'Arial'
    h2.font.size = Pt(15)
    h2.font.color.rgb = DARK_NAVY
    h2.font.bold = True
    h2_pf = h2.paragraph_format
    h2_pf.space_before = Pt(14)
    h2_pf.space_after = Pt(6)

    h3 = styles['Heading 3']
    h3.font.name = 'Arial'
    h3.font.size = Pt(13)
    h3.font.color.rgb = ROYAL_BLUE
    h3.font.bold = True

    normal = styles['Normal']
    normal.font.name = 'Arial'
    normal.font.size = Pt(11)
    normal.font.color.rgb = BODY_TEXT
    normal_pf = normal.paragraph_format
    normal_pf.space_after = Pt(4)

    # ── Header and Footer ───────────────────────────────────────
    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    header = section.header
    header_para = header.paragraphs[0]
    run1 = header_para.add_run("Artisan's Echo")
    run1.font.color.rgb = SOFT_GOLD
    run1.font.size = Pt(9)
    run1.font.italic = True
    run2 = header_para.add_run("    |    Rapport PFE - AIT OUAHMAN OUISSAL")
    run2.font.color.rgb = MID_GREY
    run2.font.size = Pt(8.5)
    run2.font.italic = True
    header_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    footer = section.footer
    footer_para = footer.paragraphs[0]
    footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_f1 = footer_para.add_run("Racine - 2025/2026    |    Page ")
    run_f1.font.size = Pt(8.5)
    run_f1.font.color.rgb = MID_GREY
    add_page_number(footer_para.add_run())

    # ── Read markdown file ──────────────────────────────────────
    md_path = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                           "Rapport_PFE_Artisans_Echo_AIT_OUAHMAN_OUISSAL.md")
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    in_table = False
    table_rows = []
    in_code_block = False

    for line in lines:
        line = line.rstrip('\n')
        stripped = line.strip()

        # Handle code blocks
        if stripped.startswith('```'):
            in_code_block = not in_code_block
            if not in_code_block and stripped == '```':
                continue
            continue
        if in_code_block:
            p = doc.add_paragraph()
            run = p.add_run(line)
            run.font.name = 'Courier New'
            run.font.size = Pt(8.5)
            run.font.color.rgb = BODY_TEXT
            continue

        # Handle tables
        if stripped.startswith('|'):
            in_table = True
            row_data = [cell.strip() for cell in stripped.split('|')[1:-1]]
            if all(cell.replace('-', '').strip() == '' for cell in row_data):
                continue
            table_rows.append(row_data)
            continue
        else:
            if in_table and len(table_rows) > 0:
                num_cols = len(table_rows[0])
                table = doc.add_table(rows=len(table_rows), cols=num_cols)
                table.alignment = WD_TABLE_ALIGNMENT.CENTER

                for i, row in enumerate(table_rows):
                    for j, cell_text in enumerate(row):
                        if j < len(table.columns):
                            cell = table.cell(i, j)
                            p = cell.paragraphs[0]
                            p.space_before = Pt(2)
                            p.space_after = Pt(2)
                            parse_inline(p, cell_text)
                            set_cell_borders(cell)
                            set_cell_padding(cell)

                            if i == 0:
                                # Header row - Royal Blue background
                                set_cell_background(cell, "1E3A8A")
                                for run in p.runs:
                                    run.font.color.rgb = WHITE
                                    run.font.bold = True
                                    run.font.size = Pt(9.5)
                            else:
                                if i % 2 == 0:
                                    set_cell_background(cell, "EFF6FF")  # light blue
                                else:
                                    set_cell_background(cell, "FFFFFF")
                                for run in p.runs:
                                    run.font.size = Pt(9.5)
                                    run.font.color.rgb = BODY_TEXT

                doc.add_paragraph()
                in_table = False
                table_rows = []

        if not stripped:
            continue

        if stripped == '---':
            doc.add_page_break()
            continue

        # Screenshot placeholders
        if stripped.startswith('[Capture'):
            # Extract figure label and caption
            match = re.search(r'Figure\s+(\d+\.?\d*)\s*:\s*(.+?)\]', stripped)
            if match:
                fig_num = match.group(1)
                fig_caption = match.group(2).strip()
                create_screenshot_placeholder(doc, f"Figure {fig_num}", fig_caption)
            else:
                create_screenshot_placeholder(doc, "Figure", stripped)
            continue

        # Italic caption lines starting with *
        if stripped.startswith('*') and stripped.endswith('*') and not stripped.startswith('**'):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(stripped[1:-1])
            run.font.italic = True
            run.font.color.rgb = MID_GREY
            run.font.size = Pt(9.5)
            continue

        # Headings
        if stripped.startswith('# '):
            text = stripped[2:].strip()
            if text.startswith('Chapitre'):
                doc.add_page_break()
                # Chapter separator
                sep_p = doc.add_paragraph()
                sep_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                sep_p.space_before = Pt(80)
                sep_p.space_after = Pt(30)
                run = sep_p.add_run(text)
                run.font.size = Pt(26)
                run.font.color.rgb = ROYAL_BLUE
                run.font.bold = True
                # Gold line
                line_p = doc.add_paragraph()
                line_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = line_p.add_run("_" * 40)
                run.font.color.rgb = SOFT_GOLD
                continue
            doc.add_heading(text, level=1)
        elif stripped.startswith('## '):
            doc.add_heading(stripped[3:].strip(), level=2)
        elif stripped.startswith('### '):
            doc.add_heading(stripped[4:].strip(), level=3)
        elif stripped.startswith('#### '):
            doc.add_heading(stripped[5:].strip(), level=3)
        elif stripped.startswith('- '):
            # Bullet point
            p = doc.add_paragraph(style='List Bullet')
            parse_inline(p, stripped[2:])
        else:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            parse_inline(p, stripped)

    # ── Save DOCX ───────────────────────────────────────────────
    out_path = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                            "Rapport_PFE_Artisans_Echo_AIT_OUAHMAN_OUISSAL.docx")
    doc.save(out_path)
    print(f"DOCX cree avec succes : {out_path}")

if __name__ == "__main__":
    main()
